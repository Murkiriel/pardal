"""DNIT — PNCV "Controle de Velocidade" (rodovia federal, país todo).

Descobre o XLSX mais novo pela API do CKAN; coluna "Coordenadas (Lat/Long)" no formato
"LAT / LNG". PNCV não traz limite. source="DNIT".

"Faixas" diz o sentido de cada faixa fiscalizada ("P-C-1" = pista crescente, "P-D-1" =
decrescente). Só um dos dois -> direction_deg pela rota do SNV da BR (ver common/direction.py);
os dois -> vazio.

"Rodovia" + "Km" vão em road_km: é o km do cadastro, o mesmo que o Inmetro usa, e por ele o radar
casa com o medidor (e ganha o limite dele) em inmetro_status.py.
"""
from __future__ import annotations

import io
import re
from typing import Callable, List, Optional, Tuple

import openpyxl

from datakit.context import SourceData, BuildContext
from datakit.common import Camera, CameraKind, Limit, in_bbox
from datakit.common.lrs import SnvRoutes
from datakit.common.direction import direction_on
from datakit.sources import snv
from datakit.sources._http import ckan_resources, get_bytes, in_br, newest, to_float

CKAN = "https://servicos.dnit.gov.br/dadosabertos/api/3/action/package_show?id=controle-de-velocidade"
SNAP_M = 150.0


def road_km(rodovia, km) -> Optional[Tuple[int, float]]:
    """('070', 189.8) -> (70, 189.8); sem rodovia ou km legíveis -> None."""
    m = re.search(r"(\d{2,3})", str(rodovia or ""))
    k = to_float(km) if km is not None else None
    return (int(m.group(1)), k) if m and k is not None else None


def lanes_increasing(faixas: Optional[str]) -> Optional[bool]:
    """'P-C-1, P-C-2' -> True, 'P-D-1' -> False, 'P-C-1, P-D-1' (os dois) ou vazio -> None."""
    sides = set(re.findall(r"P-([CD])-\d", str(faixas or "").upper()))
    return (sides.pop() == "C") if len(sides) == 1 else None


def load(raw_dir: str, bbox: Optional[Tuple[float, float, float, float]] = None,
         routes: Optional[Callable[[], Optional[SnvRoutes]]] = None) -> Tuple[List[Camera], List[Limit]]:
    """`routes`: devolve as rotas do SNV (para o sentido); sem ele, carrega do raw_dir."""
    res = newest(ckan_resources(CKAN), "XLSX", "controle de velocidade")
    wb = openpyxl.load_workbook(io.BytesIO(get_bytes(res["url"])), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    try:
        ci = header.index("Coordenadas (Lat/Long)")
    except ValueError:
        raise RuntimeError(f"coluna de coordenadas não encontrada; header={header}") from None
    col = {name: header.index(name) for name in ("UF", "Rodovia", "Faixas") if name in header}
    ki = header.index("Km") if "Km" in header else None
    get_routes = routes or (lambda: snv.load_routes(raw_dir))
    snv_routes = get_routes() if len(col) == 3 else None

    out: List[Camera] = []
    for row in rows:
        cell = row[ci] if ci < len(row) else None
        if not cell or "/" not in str(cell):
            continue
        a, _, b = str(cell).partition("/")
        lat, lng = to_float(a), to_float(b)
        if lat is None or lng is None or not in_br(lat, lng) or not in_bbox(bbox, lat, lng):
            continue
        direction = _direction(snv_routes, row, col, lat, lng) if snv_routes is not None else None
        rk = road_km(row[col["Rodovia"]], row[ki]) if "Rodovia" in col and ki is not None else None
        out.append(Camera(round(lat, 6), round(lng, 6), CameraKind.FIXED, None, "DNIT", True,
                          direction_deg=direction, road_km=rk))
    return out, []


def _direction(routes, row, col, lat: float, lng: float) -> Optional[int]:
    increasing = lanes_increasing(row[col["Faixas"]])
    m = re.search(r"(\d{2,3})", str(row[col["Rodovia"]] or ""))
    if increasing is None or m is None:
        return None
    hit = routes.nearest(int(m.group(1)), str(row[col["UF"]] or "").strip(), (lat, lng), SNAP_M)
    if hit is None:
        return None
    line, km, _ = hit
    return direction_on(line, km, increasing)


def fetch(ctx: BuildContext) -> SourceData:
    """Contrato das fontes (datakit/context.py); as rotas do SNV vêm do contexto (uma carga só)."""
    return SourceData(*load(ctx.raw_dir, routes=ctx.snv_routes))
